"""Выгрузка отчётов в XLSX и CSV: файл собирается во временной директории
и отправляется вложением в личный диалог.
"""

from __future__ import annotations

import csv
import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.report_data import load_chat_report, load_staff_report, load_summary

_EXPORT_DIR = Path("/tmp/exports")

_CHAT_HEADERS = ["Чат", "Входящих", "Исходящих", "Символы клиентов", "Символы компании", "Последняя активность"]
_STAFF_HEADERS = ["Сотрудник", "Сообщений", "Символов", "Чатов", "Активных дней"]


def _fmt_dt(value: datetime | None) -> str:
    # В рабочем поясе — как в боте.
    if not value:
        return ""
    from zoneinfo import ZoneInfo

    from app.config import get_settings

    return value.astimezone(ZoneInfo(get_settings().tz)).strftime("%d.%m.%Y %H:%M")


async def build_export(
    session: AsyncSession,
    *,
    scope: str,
    target_id: int,
    start: datetime,
    end: datetime,
    period_label: str,
    fmt: str,
) -> tuple[Path, str, str] | None:
    """Собрать файл: (путь на диске, подпись, имя для человека) или None, если данных нет.
    Имя на диске уникально, чтобы одновременные запросы не затирали друг друга.
    """
    _EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if scope == "all":
        data = await load_summary(session, start, end)
        chat_rows = [
            [
                item["title"] or "",
                item["incoming"],
                item["outgoing"],
                item["client_chars"],
                item["company_chars"],
                _fmt_dt(item["last_activity"]),
            ]
            for item in data["chats"]
        ]
        staff_rows = [
            [
                item["full_name"],
                item["messages"],
                item["chars"],
                item["chats_touched"],
                item["active_days"],
            ]
            for item in data["staff"]
        ]
        sections = [("Чаты", _CHAT_HEADERS, chat_rows), ("Сотрудники", _STAFF_HEADERS, staff_rows)]
        name = f"summary_{stamp}"
        caption = f"Сводный отчёт · {period_label}"

    elif scope == "chat":
        data = await load_chat_report(session, target_id, start, end)
        if data is None:
            return None
        staff_rows = [[i["full_name"], i["messages"], i["chars"]] for i in data["staff"]]
        sections = [
            (
                "Итоги",
                ["Показатель", "Значение"],
                [
                    ["Чат", data["title"]],
                    ["Входящих", data["incoming"]],
                    ["Исходящих", data["outgoing"]],
                    ["Символы клиентов", data["client_chars"]],
                    ["Символы компании", data["company_chars"]],
                ],
            ),
            ("Сотрудники в чате", ["Сотрудник", "Сообщений", "Символов"], staff_rows),
        ]
        name = f"chat_{target_id}_{stamp}"
        caption = f"Отчёт по чату «{data['title']}» · {period_label}"

    elif scope == "staff":
        data = await load_staff_report(session, target_id, start, end)
        if data is None:
            return None
        chat_rows = [
            [i["title"] or "", i["messages"], i["chars"], _fmt_dt(i["last_activity"])]
            for i in data["chats"]
        ]
        sections = [
            (
                "Итоги",
                ["Показатель", "Значение"],
                [
                    ["Сотрудник", data["full_name"]],
                    ["Сообщений", data["messages"]],
                    ["Символов", data["chars"]],
                    ["Активных дней", data["active_days"]],
                ],
            ),
            ("По чатам", ["Чат", "Сообщений", "Символов", "Последняя активность"], chat_rows),
        ]
        name = f"staff_{target_id}_{stamp}"
        caption = f"Отчёт по сотруднику {data['full_name']} · {period_label}"

    else:
        return None

    # Уникальное имя на диске: иначе `path.unlink()` одного запроса унёс бы файл
    # другого. Человеку показывается `filename`.
    suffix = ".xlsx" if fmt == "xlsx" else ".csv"
    handle, raw_path = tempfile.mkstemp(prefix=f"{name}_", suffix=suffix, dir=_EXPORT_DIR)
    os.close(handle)
    path = Path(raw_path)

    if fmt == "xlsx":
        _write_xlsx(path, sections, period_label)
    else:
        _write_csv(path, sections, period_label)

    return path, caption, f"{name}{suffix}"


# Символы, с которых Excel и LibreOffice начинают разбирать ячейку как формулу.
_FORMULA_STARTERS = ("=", "+", "-", "@", "\t", "\r", "\n")


def _defuse(value):
    """Обезвредить значение, похожее на формулу (названия чатов задают участники):
    ведущий апостроф Excel понимает как «это текст».
    """
    if not isinstance(value, str):
        return value
    # Проверяются обе формы — исходная и с обрезанными ведущими пробельными символами: часть импортёров
    # пробелы обрезает, а «\tтекст» должен обезвреживаться и так.
    if value.startswith(_FORMULA_STARTERS) or value.lstrip().startswith(_FORMULA_STARTERS):
        return "'" + value
    return value


def _defuse_rows(rows):
    return [[_defuse(cell) for cell in row] for row in rows]


def _write_xlsx(path: Path, sections, period_label: str) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    bold = Font(bold=True)
    for title, headers, rows in sections:
        sheet = workbook.create_sheet(title[:31])
        sheet.append([f"Период: {period_label}"])
        sheet["A1"].font = bold
        sheet.append(headers)
        for cell in sheet[2]:
            cell.font = bold
        for row in _defuse_rows(rows):
            sheet.append(row)
        for index, header in enumerate(headers, start=1):
            column = sheet.cell(row=2, column=index).column_letter
            width = max(len(str(header)), *(len(str(r[index - 1])) for r in rows)) if rows else len(str(header))
            sheet.column_dimensions[column].width = min(width + 2, 60)
    workbook.save(path)


def _write_csv(path: Path, sections, period_label: str) -> None:
    # utf-8-sig: иначе Excel на Windows искажает кириллицу.
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow([f"Период: {period_label}"])
        for title, headers, rows in sections:
            writer.writerow([])
            writer.writerow([title])
            writer.writerow(headers)
            writer.writerows(_defuse_rows(rows))
